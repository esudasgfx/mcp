"""Excel workbook reader and Gemini analyzer tool."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

import openpyxl

from core.base_tool import BaseTool
from core.gemini_client import GeminiClient


class ExcelTool(BaseTool):
    """Read workbook rows and ask Gemini to analyze them."""

    def __init__(self) -> None:
        self._gemini = GeminiClient()

    def get_name(self) -> str:
        return "excel_tool"

    def get_description(self) -> str:
        return "Reads an Excel workbook and uses Gemini to summarize or analyze it."

    def get_input_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_path": {
                    "type": "string",
                    "description": "Path to the .xlsx workbook on the MCP server machine.",
                },
                "query": {
                    "type": "string",
                    "description": "Question or instruction for Gemini.",
                },
                "worksheet": {
                    "type": "string",
                    "description": "Optional worksheet name. Defaults to the active sheet.",
                },
                "max_rows": {
                    "type": "integer",
                    "description": "Maximum rows to read from the worksheet.",
                    "default": 200,
                    "minimum": 1,
                    "maximum": 5000,
                },
            },
            "required": ["file_path", "query"],
            "additionalProperties": False,
        }

    async def execute(self, params: dict[str, Any]) -> str:
        file_path = Path(str(params["file_path"])).expanduser()
        query = str(params["query"])
        worksheet_name = params.get("worksheet")
        max_rows = int(params.get("max_rows", 200))

        if not file_path.exists():
            return f"Excel file not found: {file_path}"

        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            worksheet = (
                workbook[str(worksheet_name)]
                if worksheet_name
                else workbook.active
            )
            buffer = io.StringIO()
            writer = csv.writer(buffer)

            for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if index > max_rows:
                    break
                writer.writerow(["" if value is None else value for value in row])

            raw_data = (
                f"Workbook: {file_path.name}\n"
                f"Worksheet: {worksheet.title}\n"
                f"Rows sampled: {min(worksheet.max_row, max_rows)}\n\n"
                f"{buffer.getvalue()}"
            )
        finally:
            workbook.close()

        return await self._gemini.process_data(raw_data, query)
